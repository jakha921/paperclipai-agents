from __future__ import annotations

import io
from typing import Any

from django.db.models import Count

from .models import Department, Employee


class OrgChartService:
    """Сервис для получения данных org chart."""

    @staticmethod
    def get_org_chart_data() -> list[dict[str, Any]]:
        nodes = (
            Department.objects.filter(is_active=True)
            .select_related("head")
            .annotate(employee_count=Count("employees"))
        )
        return [
            {
                "id": str(dept.id),
                "name": dept.get_name(),
                "parentId": (str(dept.parent_id) if dept.parent_id else None),
                "employee_count": dept.employee_count,
                "head_name": (dept.head.full_name if dept.head else None),
            }
            for dept in nodes
        ]


class ExcelExportService:
    """Экспорт сотрудников в Excel."""

    @staticmethod
    def export_employees(queryset=None) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Сотрудники"

        headers = [
            "Табельный номер",
            "Фамилия",
            "Имя",
            "Отчество",
            "Подразделение",
            "Должность",
            "Категория",
            "Дата найма",
            "Статус",
            "Телефон",
            "Email",
            "ПИНФЛ",
        ]
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill

        if queryset is None:
            queryset = Employee.objects.filter(is_deleted=False).select_related(
                "department", "position"
            )

        for row, emp in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=emp.employee_number)
            ws.cell(row=row, column=2, value=emp.last_name)
            ws.cell(row=row, column=3, value=emp.first_name)
            ws.cell(row=row, column=4, value=emp.middle_name)
            ws.cell(row=row, column=5, value=emp.department.get_name())
            ws.cell(row=row, column=6, value=emp.position.name.get("ru", ""))
            ws.cell(row=row, column=7, value=emp.position.category)
            ws.cell(row=row, column=8, value=str(emp.hire_date))
            ws.cell(row=row, column=9, value=emp.status)
            ws.cell(row=row, column=10, value=emp.phone)
            ws.cell(row=row, column=11, value=emp.email)
            ws.cell(row=row, column=12, value=emp.pinfl)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
