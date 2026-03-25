import io
import uuid
from datetime import date

from django.core.files.base import ContentFile

from .models import DocumentTemplate, GeneratedDocument


def render_template(template: DocumentTemplate, context: dict) -> str:
    """Рендер Jinja2 шаблона."""
    from jinja2 import Template as Jinja2Template

    jinja_template = Jinja2Template(template.template_html)
    return jinja_template.render(**context)


def generate_pdf(html_content: str) -> bytes:
    """HTML -> PDF bytes через WeasyPrint."""
    from weasyprint import HTML

    return HTML(string=html_content).write_pdf()


def generate_document(
    template_code: str,
    employee_id: str,
    extra_data: dict,
    user,
) -> GeneratedDocument:
    """Создать GeneratedDocument, сохранить PDF файл."""
    from apps.departments.models import Employee

    template = DocumentTemplate.objects.get(code=template_code, is_active=True)
    employee = Employee.objects.select_related("department", "position").get(id=employee_id)

    context = {
        "employee": employee,
        "date": date.today(),
        "order_number": str(uuid.uuid4())[:8].upper(),
        **extra_data,
    }

    html_content = render_template(template, context)
    pdf_bytes = generate_pdf(html_content)

    doc = GeneratedDocument(
        template=template,
        employee=employee,
        generated_by=user,
        data={"extra_data": extra_data},
    )
    filename = f"{template.code}_{employee.id}_{date.today()}.pdf"
    doc.file.save(filename, ContentFile(pdf_bytes), save=True)

    return doc


def generate_payroll_excel(period_year: int, period_month: int, department_id: str = None) -> bytes:
    """Excel ведомость за период (год + месяц)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    from apps.payroll.models import Payroll

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll"

    headers = [
        "#",
        "ФИО",
        "Отдел",
        "Должность",
        "Оклад",
        "Надбавки",
        "Удержания",
        "К выплате",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    qs = Payroll.objects.filter(year=period_year, month=period_month).select_related(
        "employee__department", "employee__position"
    )
    if department_id:
        qs = qs.filter(employee__department_id=department_id)

    for row_num, run in enumerate(qs, 2):
        bonuses = float(
            run.academic_bonus + run.position_bonus + run.seniority_bonus + run.other_allowances
        )
        ws.cell(row=row_num, column=1, value=row_num - 1)
        ws.cell(row=row_num, column=2, value=str(run.employee))
        ws.cell(row=row_num, column=3, value=str(run.employee.department))
        ws.cell(row=row_num, column=4, value=str(run.employee.position))
        ws.cell(row=row_num, column=5, value=float(run.base_salary))
        ws.cell(row=row_num, column=6, value=bonuses)
        ws.cell(row=row_num, column=7, value=float(run.total_deductions))
        ws.cell(row=row_num, column=8, value=float(run.net_salary))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def generate_employee_excel(filters: dict) -> bytes:
    """Excel список сотрудников."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    from apps.departments.models import Employee

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"

    headers = ["#", "ФИО", "Email", "Отдел", "Должность", "Дата найма", "Статус"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    qs = Employee.objects.select_related("department", "position").all()
    if filters.get("department_id"):
        qs = qs.filter(department_id=filters["department_id"])
    if filters.get("status"):
        qs = qs.filter(status=filters["status"])

    for row_num, emp in enumerate(qs, 2):
        ws.cell(row=row_num, column=1, value=row_num - 1)
        ws.cell(row=row_num, column=2, value=emp.full_name)
        ws.cell(row=row_num, column=3, value=emp.email)
        ws.cell(row=row_num, column=4, value=str(emp.department))
        ws.cell(row=row_num, column=5, value=str(emp.position))
        ws.cell(row=row_num, column=6, value=str(emp.hire_date))
        ws.cell(row=row_num, column=7, value=emp.status)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
